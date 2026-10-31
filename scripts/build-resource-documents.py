"""Build actual SSA PDFs and editable DOCX forms from the resource catalog."""
import argparse, hashlib, json, os, re, urllib.request
from pathlib import Path
from xml.sax.saxutils import escape
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader

NAVY = colors.HexColor('#24476a')
GRAY = colors.HexColor('#eff2f5')
LINE = colors.HexColor('#cbd4de')
ROOT = Path(__file__).resolve().parents[1]
FONT_NAME = 'SSA'
FONT_BOLD = 'SSA-Bold'

def fonts():
    cache = ROOT / '.document-fonts'
    cache.mkdir(exist_ok=True)
    for weight, filename, windows in [('Regular','NanumGothic-Regular.ttf','malgun.ttf'),('Bold','NanumGothic-Bold.ttf','malgunbd.ttf')]:
        local = Path('C:/Windows/Fonts') / windows
        target = local if local.exists() else cache / filename
        if not target.exists():
            urllib.request.urlretrieve('https://raw.githubusercontent.com/google/fonts/main/ofl/nanumgothic/' + filename, target)
        pdfmetrics.registerFont(TTFont(FONT_NAME if weight == 'Regular' else FONT_BOLD, str(target)))
    pdfmetrics.registerFontFamily(FONT_NAME, normal=FONT_NAME, bold=FONT_BOLD)

def catalog():
    text = (ROOT / 'app/resources-data.ts').read_text(encoding='utf-8')
    return json.loads(text.split('export const resources:Resource[]=',1)[1].strip().rstrip(';'))

def style(size=10, bold=False, **kw):
    return ParagraphStyle('ssa',fontName=FONT_BOLD if bold else FONT_NAME,fontSize=size,leading=size*1.5,wordWrap='CJK',textColor=colors.HexColor('#24394d'),spaceAfter=5,**kw)

def p(text, size=10, bold=False):
    return Paragraph(escape(text).replace('\n','<br/>'), style(size,bold))

def table(rows, widths, header=False, heights=None):
    t = Table([[p(str(v),9) for v in row] for row in rows],colWidths=widths,rowHeights=heights,hAlign='LEFT')
    t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,LINE),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),10),('RIGHTPADDING',(0,0),(-1,-1),10),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),('BACKGROUND',(0,0),(0,-1),GRAY)]))
    return t

class NumberedCanvas(canvas.Canvas):
    def __init__(self,*args,**kwargs):
        super().__init__(*args,**kwargs)
        self.pages = []
    def showPage(self):
        self.pages.append(dict(self.__dict__))
        self._startPage()
    def save(self):
        count = len(self.pages)
        for state in self.pages:
            self.__dict__.update(state)
            self.setStrokeColor(NAVY)
            self.setLineWidth(.8)
            self.line(48, A4[1]-48, A4[0]-48, A4[1]-48)
            self.setFont(FONT_BOLD,12)
            self.setFillColor(NAVY)
            self.drawString(48,A4[1]-37,'SSA')
            self.setFont(FONT_NAME,8)
            self.drawString(82,A4[1]-36,'특수치안관리기구 서울지국')
            self.setFont(FONT_NAME,7)
            self.setFillColor(colors.HexColor('#748292'))
            self.drawString(48,35,'SSA HERO NETWORK · UNKNOWN 세계관 창작 자료')
            self.drawRightString(A4[0]-48,35,str(self._pageNumber)+' / '+str(count))
            super().showPage()
        super().save()

def pdf(r,path):
    doc = SimpleDocTemplate(str(path),pagesize=A4,rightMargin=48,leftMargin=48,topMargin=68,bottomMargin=56,title=r['title'],author='SSA 서울지국 '+r['hero'])
    w = A4[0]-96
    story = [p(r['board']+' / '+r['kind'],9),p(r['title'],18,True),Spacer(1,9)]
    story.append(table([['문서번호',r['code'],'개정일',r['created']],['담당 부서',r['hero'],'판 / 구분',r['version']+' / 공통 업무자료']],[70,w/2-70,70,w/2-70]))
    story.append(Spacer(1,8))
    if r['kind']=='서식':
        story.append(table([['작성','확인','접수'],['성명 / 서명:','성명 / 서명:','접수번호 / 일자:']],[w/3]*3,heights=[26,34]))
        story.append(Spacer(1,10))
    for block in r['body'].split('\n\n'):
        lines=block.splitlines()
        heading,body=lines[0],lines[1:]
        story.append(KeepTogether([p(heading,11,True),Spacer(1,2)]))
        if any('____' in line for line in body):
            rows=[];heights=[]
            for line in body:
                if '____' in line and ':' in line:
                    label=line.split(':',1)[0]
                    rows.append([label,' '])
                    heights.append(34 if any(x in label for x in ['사항','내용','경위','업무','근거','사유','결과','상황','관측','특징','요소','제안']) else 28)
                elif line.strip():
                    if rows:
                        story.append(table(rows,[155,w-155],heights=heights));rows=[];heights=[]
                        story.append(Spacer(1,7))
                    story.append(p(line))
            if rows: story.append(table(rows,[155,w-155],heights=heights))
            story.append(Spacer(1,12))
        elif any(line.startswith('□') for line in body):
            rows=[]
            for line in body:
                if line.startswith('□'): rows.append(['□',line[1:].strip()])
                else: story.append(p(line))
            if rows: story.append(table(rows,[30,w-30]))
            story.append(Spacer(1,12))
        else:
            for line in body: story.append(p(line))
            story.append(Spacer(1,5))
    doc.build(story,canvasmaker=NumberedCanvas)

def set_font(run,size=10,bold=False):
    run.font.name='맑은 고딕';run.font.size=Pt(size);run.bold=bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')

def word_table(doc,rows,blank=False):
    t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Table Grid'
    for vals in rows:
        cells=t.add_row().cells
        for i,v in enumerate(vals):
            para=cells[i].paragraphs[0]
            set_font(para.add_run(str(v)))
            if blank and i==1: para.paragraph_format.space_after=Pt(24)
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EFF2F5');cells[0]._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph()

def docx(r,path):
    d=Document();sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(1.8)
    normal=d.styles['Normal'];normal.font.name='맑은 고딕';normal.font.size=Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    head=sec.header.paragraphs[0];set_font(head.add_run('SSA  |  특수치안관리기구 서울지국'),10,True)
    set_font(d.add_paragraph().add_run(r['title']),18,True)
    word_table(d,[['문서번호',r['code']],['자료 구분',r['board']+' / '+r['kind']],['담당 부서',r['hero']],['개정일 / 판',r['created']+' / '+r['version']]])
    word_table(d,[['작성','확인','접수'],['성명 / 서명:','성명 / 서명:','접수번호 / 일자:']])
    for block in r['body'].split('\n\n'):
        lines=block.splitlines()
        h=d.add_paragraph();h.paragraph_format.keep_with_next=True;set_font(h.add_run(lines[0]),12,True)
        fields=[]
        for line in lines[1:]:
            if '____' in line and ':' in line: fields.append([line.split(':',1)[0],''])
            elif line.startswith('□'): fields.append([line[1:].strip(),'□'])
            else:
                if fields: word_table(d,fields,True);fields=[]
                set_font(d.add_paragraph().add_run(line))
        if fields: word_table(d,fields,True)
    foot=sec.footer.paragraphs[0];set_font(foot.add_run('SSA HERO NETWORK · UNKNOWN 세계관 창작 자료  |  '),8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    d.core_properties.title=r['title'];d.core_properties.author='SSA 서울지국'
    d.save(path)

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--output',default='public/resources');args=ap.parse_args()
    output=Path(args.output);output.mkdir(parents=True,exist_ok=True);fonts();records=[]
    docs=catalog()
    assert docs and len({r['code'] for r in docs}) == len(docs)
    assert len({r['id'] for r in docs}) == len(docs)
    for r in docs:
        files=[]
        path=output/(r['code']+'.pdf');pdf(r,path)
        reader=PdfReader(path)
        text=''.join(page.extract_text() or '' for page in reader.pages)
        assert r['code'] in text and re.sub(r'\s+','',r['title']) in re.sub(r'\s+','',text)
        assert all(page.mediabox.width>590 and page.mediabox.height>840 for page in reader.pages)
        files.append(path)
        if r['kind']=='서식':
            path=output/(r['code']+'.docx');docx(r,path);check=Document(path)
            assert any(r['title'] in p.text for p in check.paragraphs) and len(check.tables)>=3
            files.append(path)
        records.append({'id':r['id'],'code':r['code'],'title':r['title'],'pages':len(reader.pages),'files':[{'name':f.name,'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in files]})
    (output/'manifest.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'pdf':len(records),'docx':sum(r['kind']=='서식' for r in docs),'validated':True,'pages':[r['pages'] for r in records]}))
if __name__=='__main__': main()
